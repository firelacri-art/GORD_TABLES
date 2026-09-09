#!/usr/bin/env python3
"""
verify_gord_table.py — проверка .xlsx на соответствие стилю GORD.

    python3 verify_gord_table.py file.xlsx [--render]

Проверяет на каждом листе: есть логотип, шрифт Geologica у всех заполненных
ячеек, серая шапка F3F3F3, красная линейка/рамка D72410, заголовок листа.
С флагом --render дополнительно конвертирует книгу в PDF через LibreOffice
(если установлен) и кладёт рядом PNG первых страниц для визуальной проверки.
Код выхода 1, если найдены нарушения.
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl

FONT = "Geologica"
RED = "FFD72410"
GREY_HEADER = "FFF3F3F3"


def check(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path)
    problems: list[str] = []
    for ws in wb.worksheets:
        tag = f"[{ws.title}]"
        if not getattr(ws, "_images", []):
            problems.append(f"{tag} нет логотипа (картинки на листе)")
        bad_fonts = set()
        has_grey_header = False
        has_red = False
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and c.font and c.font.name != FONT:
                    bad_fonts.add(c.font.name)
                if c.fill and c.fill.fill_type == "solid" and c.fill.fgColor.rgb == GREY_HEADER:
                    has_grey_header = True
                b = c.border
                for s in (b.left, b.right, b.top, b.bottom):
                    if s is not None and s.color is not None and s.color.rgb == RED:
                        has_red = True
        if bad_fonts:
            problems.append(f"{tag} шрифт не {FONT}: {', '.join(sorted(map(str, bad_fonts)))}")
        if not has_red:
            problems.append(f"{tag} нет красного акцента D72410 (линейка под логотипом / рамка названия)")
        # у листов-таблиц должна быть серая шапка; у карточек — рамка названия
        if not has_grey_header and ws.max_column > 4:
            problems.append(f"{tag} нет серой шапки таблицы F3F3F3")
        title_cells = [ws.cell(row=r, column=2).value for r in (2, 3)]
        if not any(isinstance(v, str) and v.strip() for v in title_cells):
            problems.append(f"{tag} пустой заголовок листа (B2/B3)")
    return problems


def render(path: Path) -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice не найден — визуальный рендер пропущен")
        return
    outdir = path.parent / "render"
    outdir.mkdir(exist_ok=True)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
                   check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    pdf = outdir / (path.stem + ".pdf")
    if not pdf.exists():
        print("PDF не получился")
        return
    try:
        import pypdfium2 as pdfium  # type: ignore
    except ImportError:
        print(f"PDF: {pdf} (для PNG установи pypdfium2: pip install pypdfium2)")
        return
    doc = pdfium.PdfDocument(str(pdf))
    for i in range(min(len(doc), 12)):
        png = outdir / f"{path.stem}_p{i + 1}.png"
        doc[i].render(scale=1.3).to_pil().save(png)
    print(f"PDF: {pdf}\nPNG: {outdir}/{path.stem}_p1..{min(len(doc), 12)}.png — открой и посмотри глазами")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--render", action="store_true")
    a = ap.parse_args(argv)
    path = Path(a.file)
    problems = check(path)
    wb = openpyxl.load_workbook(path)
    print(f"{path.name}: {len(wb.sheetnames)} лист(ов): {', '.join(wb.sheetnames)}")
    if problems:
        print("НАРУШЕНИЯ:")
        for p in problems:
            print(" -", p)
    else:
        print("Стиль GORD: OK (логотип, Geologica, серая шапка, красный акцент на каждом листе)")
    if a.render:
        render(path)
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
